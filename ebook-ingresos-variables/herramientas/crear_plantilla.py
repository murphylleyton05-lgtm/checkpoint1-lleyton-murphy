#!/usr/bin/env python3
"""Genera la plantilla de cálculo del Método Embalse (Excel / Google Sheets).

Produce dos archivos en entregables/:
  - Plantilla-Metodo-Embalse.xlsx          (en blanco)
  - Plantilla-Metodo-Embalse-EJEMPLO.xlsx  (con el caso hipotético de Valeria)

Las fórmulas se escriben en inglés (Excel las traduce al idioma de la
instalación) y usan solo funciones compatibles con Excel y Google Sheets.
Uso: python3 herramientas/crear_plantilla.py
"""
import datetime as dt
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

SALIDA = Path(__file__).resolve().parent.parent / "entregables"

PROFUNDO = "0E3A4A"
AGUA = "1A7FB5"
CLARO = "E9F3F9"
ENTRADA = "FFF4D6"   # celdas para completar
ARENA = "F7F3EC"

F_TIT = Font(name="Calibri", size=16, bold=True, color=PROFUNDO)
F_SUB = Font(name="Calibri", size=10, italic=True, color="5F6F77")
F_HEAD = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
F_BOLD = Font(name="Calibri", size=10, bold=True, color="15232B")
F_NORM = Font(name="Calibri", size=10, color="15232B")
F_NOTA = Font(name="Calibri", size=9, italic=True, color="5F6F77")
FILL_HEAD = PatternFill("solid", fgColor=PROFUNDO)
FILL_IN = PatternFill("solid", fgColor=ENTRADA)
FILL_OUT = PatternFill("solid", fgColor=CLARO)
FILL_ARENA = PatternFill("solid", fgColor=ARENA)
BORDE = Border(bottom=Side(style="thin", color="D8E1E5"))
NUM = '#,##0'
NUM2 = '#,##0.00'
PCT = '0.0%'

MESES = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio", "Agosto",
         "Septiembre", "Octubre", "Noviembre", "Diciembre"]


def titulo(ws, texto, sub):
    ws["A1"] = texto
    ws["A1"].font = F_TIT
    ws["A2"] = sub
    ws["A2"].font = F_SUB
    ws.sheet_view.showGridLines = False


def encabezado(ws, fila, col, textos, anchos=None):
    for i, t in enumerate(textos):
        c = ws.cell(row=fila, column=col + i, value=t)
        c.font = F_HEAD
        c.fill = FILL_HEAD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[fila].height = 30
    if anchos:
        for i, a in enumerate(anchos):
            ws.column_dimensions[chr(64 + col + i)].width = a


def celda(ws, ref, valor=None, entrada=False, fmt=None, negrita=False):
    c = ws[ref]
    if valor is not None:
        c.value = valor
    c.font = F_BOLD if negrita else F_NORM
    c.fill = FILL_IN if entrada else FILL_OUT
    c.border = BORDE
    if fmt:
        c.number_format = fmt
    return c


def etiqueta(ws, ref, texto, negrita=False):
    c = ws[ref]
    c.value = texto
    c.font = F_BOLD if negrita else F_NORM
    c.alignment = Alignment(wrap_text=True, vertical="center")
    c.border = BORDE


def lista(ws, rango, opciones):
    dv = DataValidation(type="list", formula1='"' + ",".join(opciones) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(rango)


# ------------------------------------------------------------------ hojas
def hoja_inicio(wb):
    ws = wb.active
    ws.title = "Inicio"
    titulo(ws, "Método Embalse · Plantilla de cálculo",
           "Complemento del libro «Ingresos variables, sueldo fijo». Material educativo; no constituye asesoramiento financiero ni fiscal.")
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 90
    filas = [
        ("Cómo se usa", ""),
        ("Celdas amarillas", "Son las que completas tú. Todo lo demás se calcula solo."),
        ("Celdas celestes", "Resultados calculados. No hace falta tocarlas."),
        ("", ""),
        ("Orden recomendado", ""),
        ("1 Ingresos", "Tus cobros brutos de los últimos 12 meses y tus clientes (capítulo 2)."),
        ("2 Gastos", "Tu inventario de gastos personales clasificado (capítulo 3)."),
        ("3 Sistema", "Tu porcentaje de impuestos, gastos de la actividad y respuestas para la meta del embalse. Aquí está la ficha completa de tu sistema (capítulos 4 a 7)."),
        ("6 Previsiones", "Tus gastos irregulares del año: calcula el aporte mensual y el capital semilla (capítulo 8)."),
        ("7 Tarifa", "Tu tarifa mínima viable y tu tarifa objetivo (capítulo 11)."),
        ("4 Cobros", "Uso semanal: registra cada cobro y la reserva de impuestos que le corresponde."),
        ("5 Embalse", "Uso mensual: el día de pago, anota el saldo del embalse y obtén el sueldo sugerido, la zona del semáforo y la cascada (capítulos 6, 7 y 9)."),
        ("8 Cartera 90 días", "Uso mensual: lo que esperas cobrar en los próximos tres meses, ponderado (capítulo 11)."),
        ("", ""),
        ("Google Sheets", "Archivo > Importar > Subir este archivo. Las fórmulas son compatibles."),
        ("Moneda", "Usa una sola moneda de referencia en toda la plantilla (capítulo 2)."),
        ("Aviso", "Los porcentajes de impuestos deben confirmarse con la autoridad tributaria de tu país o con un profesional. Los rangos de variabilidad y la tabla de meses del embalse son referencias propias del método."),
    ]
    for i, (a, b) in enumerate(filas, start=4):
        ws.cell(row=i, column=2, value=a).font = F_BOLD if a and not b else F_BOLD
        ws.cell(row=i, column=3, value=b).font = F_NORM
        ws.cell(row=i, column=3).alignment = Alignment(wrap_text=True, vertical="top")
        if a == "Celdas amarillas":
            ws.cell(row=i, column=2).fill = FILL_IN
        if a == "Celdas celestes":
            ws.cell(row=i, column=2).fill = FILL_OUT
        if a in ("Cómo se usa", "Orden recomendado"):
            ws.cell(row=i, column=2).font = Font(name="Calibri", size=12, bold=True, color=AGUA)


def hoja_ingresos(wb):
    ws = wb.create_sheet("1 Ingresos")
    titulo(ws, "1 · Radiografía de tus ingresos", "Registra lo que efectivamente cobraste cada mes (criterio de caja), en bruto y en tu moneda de referencia.")
    encabezado(ws, 4, 1, ["Mes", "Ingreso bruto", "Extraordinarios (no se cuentan)", "Impuestos estimados", "Gastos de la actividad", "Ingreso disponible"],
               [14, 15, 17, 15, 15, 16])
    for i in range(12):
        f = 5 + i
        celda(ws, f"A{f}", MESES[i], entrada=True)
        celda(ws, f"B{f}", entrada=True, fmt=NUM)
        celda(ws, f"C{f}", entrada=True, fmt=NUM)
        celda(ws, f"D{f}", f'=IF(B{f}="","",B{f}*\'3 Sistema\'!$C$6)', fmt=NUM)
        celda(ws, f"E{f}", f'=IF(B{f}="","",\'3 Sistema\'!$C$8)', fmt=NUM)
        celda(ws, f"F{f}", f'=IF(B{f}="","",B{f}-D{f}-E{f})', fmt=NUM)
    etiqueta(ws, "A17", "Total", negrita=True)
    for col in "BCDEF":
        celda(ws, f"{col}17", f"=SUM({col}5:{col}16)", fmt=NUM, negrita=True)
    ws["A18"] = "La columna E toma por defecto tus gastos mensuales de la actividad (hoja 3). Puedes escribir encima el valor real de cada mes."
    ws["A18"].font = F_NOTA

    ws.column_dimensions["G"].width = 3
    ws.column_dimensions["H"].width = 36
    ws.column_dimensions["I"].width = 16
    encabezado(ws, 4, 8, ["Tus cinco números y tu variabilidad", "Valor"])
    stats = [
        ("Promedio bruto", "=IFERROR(AVERAGE(B5:B16),0)", NUM),
        ("Promedio disponible", "=IFERROR(AVERAGE(F5:F16),0)", NUM),
        ("Mediana disponible", "=IFERROR(MEDIAN(F5:F16),0)", NUM),
        ("Mínimo disponible", "=IFERROR(MIN(F5:F16),0)", NUM),
        ("Piso (promedio de los 3 peores meses)", "=IFERROR(AVERAGE(SMALL(F5:F16,1),SMALL(F5:F16,2),SMALL(F5:F16,3)),0)", NUM),
        ("Máximo disponible", "=IFERROR(MAX(F5:F16),0)", NUM),
        ("Desvío estándar del ingreso bruto", "=IFERROR(STDEVP(B5:B16),0)", NUM),
        ("Coeficiente de variación (CV)", "=IFERROR(I11/I5,0)", PCT),
        ("Variabilidad", '=IF(COUNT(B5:B16)=0,"",IF(I12<0.15,"Baja",IF(I12<0.3,"Moderada",IF(I12<0.5,"Alta","Muy alta"))))', None),
        ("Atajo: piso bruto ÷ promedio bruto", "=IFERROR(AVERAGE(SMALL(B5:B16,1),SMALL(B5:B16,2),SMALL(B5:B16,3))/I5,0)", '0.00'),
        ("Meses por debajo del promedio disponible", '=COUNTIF(F5:F16,"<"&I6)', '0'),
        ("Meses con datos", "=COUNT(B5:B16)", '0'),
    ]
    for i, (t, f, fmt) in enumerate(stats, start=5):
        etiqueta(ws, f"H{i}", t)
        celda(ws, f"I{i}", f, fmt=fmt, negrita=(t in ("Piso (promedio de los 3 peores meses)", "Variabilidad")))

    # Concentración de clientes
    ws["A21"] = "Concentración de clientes (últimos 12 meses)"
    ws["A21"].font = Font(name="Calibri", size=12, bold=True, color=AGUA)
    encabezado(ws, 22, 1, ["Cliente", "Cobrado en el año", "% del total"])
    for f in range(23, 33):
        celda(ws, f"A{f}", entrada=True)
        celda(ws, f"B{f}", entrada=True, fmt=NUM)
        celda(ws, f"C{f}", f'=IF(B{f}="","",B{f}/SUM($B$23:$B$32))', fmt=PCT)
    etiqueta(ws, "H22", "Cliente principal (% del total)")
    celda(ws, "I22", "=IFERROR(MAX(B23:B32)/SUM(B23:B32),0)", fmt=PCT, negrita=True)
    etiqueta(ws, "H23", "Tres clientes principales (% del total)")
    celda(ws, "I23", "=IFERROR((LARGE(B23:B32,1)+IFERROR(LARGE(B23:B32,2),0)+IFERROR(LARGE(B23:B32,3),0))/SUM(B23:B32),0)", fmt=PCT)
    etiqueta(ws, "H24", "Señal de alerta")
    celda(ws, "I24", '=IF(I22>0.4,"Un cliente supera el 40 %",IF(I23>0.8,"3 clientes superan el 80 %","Sin alerta"))')
    ws.freeze_panes = "B5"
    return ws


def hoja_gastos(wb):
    ws = wb.create_sheet("2 Gastos")
    titulo(ws, "2 · Radiografía de tus gastos", "Gastos personales mensuales. Los gastos de tu actividad NO van aquí (van en la hoja 3).")
    encabezado(ws, 4, 1, ["Gasto", "Categoría", "Monto mensual", "¿Compromiso fijo?", "Nota"], [38, 14, 15, 14, 34])
    for f in range(5, 45):
        celda(ws, f"A{f}", entrada=True)
        celda(ws, f"B{f}", entrada=True)
        celda(ws, f"C{f}", entrada=True, fmt=NUM)
        celda(ws, f"D{f}", entrada=True)
        celda(ws, f"E{f}", entrada=True)
    # fila vinculada a previsiones
    ws["A5"] = "Gastos irregulares del año (desde la hoja 6)"
    ws["B5"] = "Previsión"
    ws["C5"] = "='6 Previsiones'!$C$5"
    ws["D5"] = "No"
    ws["E5"] = "Se calcula en la hoja 6. No la borres."
    for ref in ("A5", "B5", "C5", "D5", "E5"):
        ws[ref].fill = FILL_OUT
    lista(ws, "B6:B44", ["Esencial", "Importante", "Flexible", "Previsión"])
    lista(ws, "D6:D44", ["Sí", "No"])

    ws.column_dimensions["F"].width = 3
    ws.column_dimensions["G"].width = 34
    ws.column_dimensions["H"].width = 15
    encabezado(ws, 4, 7, ["Resumen", "Monto"])
    filas = [
        ("Esenciales", '=SUMIF(B5:B44,"Esencial",C5:C44)'),
        ("Importantes", '=SUMIF(B5:B44,"Importante",C5:C44)'),
        ("Flexibles", '=SUMIF(B5:B44,"Flexible",C5:C44)'),
        ("Previsiones (mensual)", '=SUMIF(B5:B44,"Previsión",C5:C44)'),
        ("", None),
        ("Costo de vida normal", "=H5+H6+H7+H8"),
        ("Modo ajuste (sin flexibles)", "=H5+H6+H8"),
        ("Número de supervivencia", "=H5"),
        ("", None),
        ("Compromisos fijos", '=SUMIF(D5:D44,"Sí",C5:C44)'),
        ("Tu piso (hoja 1)", "='1 Ingresos'!I9"),
        ("Regla del piso", '=IF(OR(H14=0,H15=0),"",IF(H14<=H15,"Cumple","No cumple: revisa compromisos fijos"))'),
    ]
    for i, (t, f) in enumerate(filas, start=5):
        if not t:
            continue
        etiqueta(ws, f"G{i}", t, negrita=t in ("Costo de vida normal", "Número de supervivencia"))
        celda(ws, f"H{i}", f, fmt=NUM, negrita=t in ("Costo de vida normal", "Número de supervivencia"))
    ws.freeze_panes = "A5"
    return ws


def hoja_sistema(wb):
    ws = wb.create_sheet("3 Sistema")
    titulo(ws, "3 · La ficha de tu sistema", "Completa las celdas amarillas. El resto se calcula a partir de las hojas 1, 2, 6 y 7.")
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 50
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 60

    def seccion(fila, texto):
        c = ws.cell(row=fila, column=2, value=texto)
        c.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        c.fill = FILL_HEAD
        ws.cell(row=fila, column=3).fill = FILL_HEAD
        ws.cell(row=fila, column=4).fill = FILL_HEAD

    def linea(fila, texto, valor, entrada=False, fmt=None, nota="", negrita=False):
        etiqueta(ws, f"B{fila}", texto, negrita=negrita)
        celda(ws, f"C{fila}", valor, entrada=entrada, fmt=fmt, negrita=negrita)
        ws[f"D{fila}"] = nota
        ws[f"D{fila}"].font = F_NOTA
        ws[f"D{fila}"].alignment = Alignment(wrap_text=True)

    seccion(4, "Reglas de entrada")
    linea(5, "Moneda de referencia", None, entrada=True, nota="Por ejemplo: USD, EUR, MXN.")
    linea(6, "% de reserva de impuestos (de cada cobro bruto)", None, entrada=True, fmt=PCT, nota="Capítulo 5. Si cobras IVA, además separas el 100 % del IVA (hoja 4).")
    linea(7, "Día de barrido / día de pago", None, entrada=True, nota="Por ejemplo: viernes / día 1.")
    linea(8, "Gastos de la actividad por mes", None, entrada=True, fmt=NUM, nota="Software, equipos (prorrateados), contador, comisiones… (capítulo 10).")
    linea(9, "Colchón operativo de la cuenta de Entrada", "=C8", fmt=NUM, nota="Un mes de gastos de la actividad.")

    seccion(11, "Sueldos")
    linea(12, "Costo de vida normal", "='2 Gastos'!H10", fmt=NUM)
    linea(13, "Sueldo meta calculado (redondeado a la centena)", "=ROUNDUP(C12,-2)", fmt=NUM)
    linea(14, "Sueldo meta ajustado a mano (opcional)", None, entrada=True, fmt=NUM, nota="Déjalo vacío para usar el calculado.")
    linea(15, "SUELDO META", '=IF(C14="",C13,C14)', fmt=NUM, negrita=True)
    linea(16, "Ingreso disponible promedio", "='1 Ingresos'!I6", fmt=NUM)
    linea(17, "80 % del ingreso disponible promedio", "=0.8*C16", fmt=NUM)
    linea(18, "Prueba del 80 %", '=IF(OR(C15=0,C16=0),"",IF(C15<=C17,"Cumple","No cumple: ver capítulo 6"))', negrita=True)
    linea(19, "Número de supervivencia", "='2 Gastos'!H12", fmt=NUM)
    linea(20, "Sueldo en modo ajuste", "='2 Gastos'!H11", fmt=NUM)
    linea(21, "Piso", "='1 Ingresos'!I9", fmt=NUM)
    linea(22, "SUELDO DE ARRANQUE", "=MIN(C15,ROUNDUP(MAX(C19,C21),-2))", fmt=NUM, negrita=True, nota="El mayor entre supervivencia y piso, redondeado, sin superar el sueldo meta.")
    linea(23, "Umbral de transición (embalse antes de pagarte)", "=2*C15", fmt=NUM, nota="Pasas al sueldo meta cuando el embalse alcanza este monto el día de pago.")

    seccion(25, "Meta del embalse")
    linea(26, "Variabilidad (hoja 1)", "='1 Ingresos'!I13")
    linea(27, "Meses base", '=IF(C26="","",IF(C26="Baja",3,IF(C26="Moderada",4,IF(C26="Alta",6,9))))', fmt='0')
    linea(28, "¿Un cliente aporta más del 40 %? (+2)", "=IF('1 Ingresos'!I22>0.4,\"Sí\",\"No\")", nota="Se calcula con la hoja 1.")
    linea(29, "¿Personas a cargo o eres el único ingreso del hogar? (+2)", None, entrada=True)
    linea(30, "¿Temporada baja con meses cercanos a cero? (+1)", None, entrada=True)
    linea(31, "¿Otro ingreso estable en el hogar cubre los esenciales? (−1)", None, entrada=True)
    linea(32, "M (meses de sueldo meta, entre 3 y 12)", '=IF(C27="","",MAX(3,MIN(12,C27+IF(C28="Sí",2,0)+IF(C29="Sí",2,0)+IF(C30="Sí",1,0)-IF(C31="Sí",1,0))))', fmt='0', negrita=True)
    linea(33, "META DEL EMBALSE", '=IF(C32="","",C32*C15)', fmt=NUM, negrita=True)
    lista(ws, "C29:C31", ["Sí", "No"])

    seccion(35, "Semáforo (sueldos meta en el embalse después de pagarte)")
    linea(36, "Verde: desde", "=3*C15", fmt=NUM, nota="3 sueldos meta o más: funcionamiento normal.")
    linea(37, "Amarillo: desde", "=1*C15", fmt=NUM, nota="Entre 1 y 3: sin compromisos fijos nuevos.")
    linea(38, "Naranja: desde", "=0.5*C15", fmt=NUM, nota="Entre medio y 1: modo ajuste y plan comercial.")
    linea(39, "Rojo: por debajo de", "=0.5*C15", fmt=NUM, nota="Modo supervivencia y protocolo rojo.")

    seccion(41, "Previsiones y tarifa")
    linea(42, "Aporte mensual a previsiones", "='6 Previsiones'!C5", fmt=NUM)
    linea(43, "Capital semilla del primer año", "='6 Previsiones'!C6", fmt=NUM)
    linea(44, "Tarifa mínima viable (por hora)", "='7 Tarifa'!C17", fmt=NUM2)
    linea(45, "Tarifa objetivo (por hora)", "='7 Tarifa'!C18", fmt=NUM2)
    return ws


def hoja_cobros(wb):
    ws = wb.create_sheet("4 Cobros")
    titulo(ws, "4 · Registro de cobros", "Una línea por cobro. Úsala en tu ritual semanal (capítulo 14).")
    encabezado(ws, 4, 1, ["Fecha", "Cliente", "Ingreso bruto (sin IVA)", "IVA cobrado", "% de reserva", "A Impuestos", "Queda en Entrada", "Mes (AAAAMM)", "¿Barrido hecho?"],
               [12, 26, 16, 12, 11, 13, 15, 12, 12])
    for f in range(5, 305):
        celda(ws, f"A{f}", entrada=True, fmt="dd/mm/yyyy")
        celda(ws, f"B{f}", entrada=True)
        celda(ws, f"C{f}", entrada=True, fmt=NUM2)
        celda(ws, f"D{f}", entrada=True, fmt=NUM2)
        celda(ws, f"E{f}", f'=IF(C{f}="","",\'3 Sistema\'!$C$6)', fmt=PCT)
        celda(ws, f"F{f}", f'=IF(C{f}="","",C{f}*E{f}+N(D{f}))', fmt=NUM2)
        celda(ws, f"G{f}", f'=IF(C{f}="","",C{f}+N(D{f})-F{f})', fmt=NUM2)
        celda(ws, f"H{f}", f'=IF(A{f}="","",YEAR(A{f})*100+MONTH(A{f}))', fmt='0')
        celda(ws, f"I{f}", entrada=True)
    lista(ws, "I5:I304", ["Sí", "No"])
    ws.column_dimensions["J"].width = 3
    ws.column_dimensions["K"].width = 30
    ws.column_dimensions["L"].width = 16
    encabezado(ws, 4, 11, ["Totales del registro", "Monto"])
    etiqueta(ws, "K5", "Ingreso bruto registrado")
    celda(ws, "L5", "=SUM(C5:C304)", fmt=NUM2)
    etiqueta(ws, "K6", "Enviado a Impuestos")
    celda(ws, "L6", "=SUM(F5:F304)", fmt=NUM2)
    etiqueta(ws, "K7", "Cobros sin barrido marcado")
    celda(ws, "L7", '=COUNTIFS(C5:C304,"<>",I5:I304,"<>Sí")', fmt='0')
    etiqueta(ws, "K9", "Bruto de un mes (escribe AAAAMM)")
    celda(ws, "L9", None, entrada=True, fmt='0')
    etiqueta(ws, "K10", "Total bruto de ese mes")
    celda(ws, "L10", '=IF(L9="","",SUMIF(H5:H304,L9,C5:C304))', fmt=NUM2)
    ws.freeze_panes = "A5"
    return ws


def hoja_embalse(wb):
    ws = wb.create_sheet("5 Embalse")
    titulo(ws, "5 · Seguimiento del embalse", "El día de pago, anota el saldo del embalse ANTES de pagarte. La hoja sugiere el sueldo, calcula la zona del semáforo y la cascada.")
    encabezado(ws, 4, 1, ["Día de pago", "Embalse antes de pagarte", "Fase", "Sueldo sugerido", "Sueldo pagado (si fue distinto)", "Embalse después", "Sueldos meta", "Zona", "Excedente sobre la meta", "Disfrute (10 %)", "Cascada (90 %)"],
               [13, 15, 11, 13, 15, 14, 10, 11, 14, 12, 12])
    SM, SA, NS, AJ, META = "'3 Sistema'!$C$15", "'3 Sistema'!$C$22", "'3 Sistema'!$C$19", "'3 Sistema'!$C$20", "'3 Sistema'!$C$33"
    for f in range(5, 41):
        celda(ws, f"A{f}", entrada=True, fmt="dd/mm/yyyy")
        celda(ws, f"B{f}", entrada=True, fmt=NUM)
        if f == 5:
            celda(ws, f"C{f}", f'=IF(B{f}="","",IF(B{f}>=2*{SM},"Normal","Arranque"))')
        else:
            celda(ws, f"C{f}", f'=IF(B{f}="","",IF(OR(C{f-1}="Normal",B{f}>=2*{SM}),"Normal","Arranque"))')
        celda(ws, f"D{f}", f'=IF(B{f}="","",IF(B{f}>=2*{SM},{SM},IF(C{f}="Arranque",MIN({SA},B{f}),IF(B{f}-{AJ}>=0.5*{SM},{AJ},MIN({NS},B{f})))))', fmt=NUM)
        celda(ws, f"E{f}", entrada=True, fmt=NUM)
        celda(ws, f"F{f}", f'=IF(B{f}="","",B{f}-IF(E{f}="",D{f},E{f}))', fmt=NUM)
        celda(ws, f"G{f}", f'=IF(B{f}="","",IFERROR(F{f}/{SM},0))', fmt='0.0')
        celda(ws, f"H{f}", f'=IF(B{f}="","",IF(G{f}>=3,"Verde",IF(G{f}>=1,"Amarillo",IF(G{f}>=0.5,"Naranja","Rojo"))))')
        celda(ws, f"I{f}", f'=IF(OR(B{f}="",N({META})=0),"",MAX(0,F{f}-{META}))', fmt=NUM)
        celda(ws, f"J{f}", f'=IF(I{f}="","",I{f}*0.1)', fmt=NUM)
        celda(ws, f"K{f}", f'=IF(I{f}="","",I{f}*0.9)', fmt=NUM)
    colores = {"Verde": "D7EED8", "Amarillo": "FBF0C2", "Naranja": "FBD9BD", "Rojo": "F6C9C9"}
    for zona, color in colores.items():
        ws.conditional_formatting.add("H5:H40", CellIsRule(operator="equal", formula=[f'"{zona}"'], fill=PatternFill("solid", fgColor=color), font=Font(bold=True)))
    ws["A42"] = ("Sueldo sugerido: sueldo meta si el embalse tiene al menos 2 sueldos meta antes de pagarte; en fase de arranque, el sueldo de arranque; "
                 "después del arranque, el sueldo en modo ajuste (zona naranja) o el número de supervivencia (zona roja). Si te pagaste otro monto, anótalo en la columna E.")
    ws["A42"].font = F_NOTA
    ws.freeze_panes = "B5"
    return ws


def hoja_previsiones(wb):
    ws = wb.create_sheet("6 Previsiones")
    titulo(ws, "6 · Gastos irregulares y previsiones", "Lista los gastos que no son mensuales, con el mes en que llegan (1 = enero … 12 = diciembre).")
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 14
    etiqueta(ws, "B4", "Total anual de gastos irregulares", negrita=True)
    celda(ws, "C4", "=SUM(D11:D40)", fmt=NUM, negrita=True)
    etiqueta(ws, "B5", "Aporte mensual a Previsiones (total ÷ 12)", negrita=True)
    celda(ws, "C5", "=C4/12", fmt=NUM, negrita=True)
    etiqueta(ws, "B6", "Capital semilla del primer año")
    celda(ws, "C6", "=MAX(0,-MIN(J11:J22))", fmt=NUM)
    etiqueta(ws, "B7", "Mes en que empiezas (1 a 12)")
    celda(ws, "C7", 1, entrada=True, fmt='0')
    encabezado(ws, 10, 2, ["Gasto", "Mes (1 a 12)", "Monto"])
    for f in range(11, 41):
        celda(ws, f"B{f}", entrada=True)
        celda(ws, f"C{f}", entrada=True, fmt='0')
        celda(ws, f"D{f}", entrada=True, fmt=NUM)
    lista(ws, "C11:C40", [str(i) for i in range(1, 13)])
    ws.column_dimensions["E"].width = 3
    for col, w in zip("FGHIJ", [12, 10, 12, 12, 14]):
        ws.column_dimensions[col].width = w
    encabezado(ws, 10, 6, ["Paso", "Mes", "Aporte", "Gastos", "Saldo simulado"])
    for k in range(12):
        f = 11 + k
        celda(ws, f"F{f}", k + 1, fmt='0')
        celda(ws, f"G{f}", f"=MOD($C$7-1+{k},12)+1", fmt='0')
        celda(ws, f"H{f}", "=$C$5", fmt=NUM)
        celda(ws, f"I{f}", f'=SUMIF($C$11:$C$40,G{f},$D$11:$D$40)', fmt=NUM)
        celda(ws, f"J{f}", f"=H{f}-I{f}" if k == 0 else f"=J{f-1}+H{f}-I{f}", fmt=NUM)
    ws["F24"] = "Si el saldo simulado es negativo en algún mes, el capital semilla es lo que necesitas poner al empezar para que nunca falte."
    ws["F24"].font = F_NOTA
    ws.conditional_formatting.add("J11:J22", CellIsRule(operator="lessThan", formula=["0"], fill=PatternFill("solid", fgColor="F6C9C9")))
    return ws


def hoja_tarifa(wb):
    ws = wb.create_sheet("7 Tarifa")
    titulo(ws, "7 · Tu tarifa mínima viable", "Capítulo 11. Cuánto necesitas cobrar por hora facturable para sostener todo el sistema.")
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 48
    ws.column_dimensions["C"].width = 16
    filas = [
        (4, "Sueldo meta × 12", "='3 Sistema'!C15*12", False, NUM),
        (5, "Gastos de la actividad × 12", "='3 Sistema'!C8*12", False, NUM),
        (6, "Metas y retiro por año", None, True, NUM),
        (7, "Ingreso neto anual necesario", "=C4+C5+N(C6)", False, NUM),
        (8, "% de reserva de impuestos", "='3 Sistema'!C6", False, PCT),
        (9, "Ingreso bruto anual necesario", "=IFERROR(C7/(1-C8),0)", False, NUM),
        (10, "Semanas de vacaciones", 4, True, '0'),
        (11, "Semanas de feriados", 2, True, '0'),
        (12, "Semanas de margen (enfermedad, imprevistos)", 1, True, '0'),
        (13, "Horas de trabajo por semana", 40, True, '0'),
        (14, "% de horas facturables", 0.55, True, PCT),
        (15, "Semanas de trabajo por año", "=52-C10-C11-C12", False, '0'),
        (16, "Horas facturables por año", "=C15*C13*C14", False, NUM),
        (17, "TARIFA MÍNIMA VIABLE (por hora)", "=IFERROR(C9/C16,0)", False, NUM2),
        (18, "TARIFA OBJETIVO (× 1,15)", "=C17*1.15", False, NUM2),
        (19, "Tu tarifa promedio actual (por hora)", None, True, NUM2),
        (20, "Diferencia con la tarifa objetivo", '=IF(C19="","",C19-C18)', False, NUM2),
    ]
    for f, t, v, e, fmt in filas:
        etiqueta(ws, f"B{f}", t, negrita=t.isupper() or t.startswith("TARIFA"))
        celda(ws, f"C{f}", v, entrada=e, fmt=fmt, negrita=t.startswith("TARIFA"))
    ws["B22"] = "Para proyectos: horas estimadas × 1,2 (imprevistos) × tarifa objetivo."
    ws["B22"].font = F_NOTA
    return ws


def hoja_cartera(wb):
    ws = wb.create_sheet("8 Cartera 90 días")
    titulo(ws, "8 · Cartera de trabajos a 90 días", "Lo que esperas cobrar en los próximos tres meses. Confirmado = 100 %, Probable = 70 %, Posible = 30 % (convención práctica).")
    encabezado(ws, 4, 1, ["Cliente / trabajo", "Monto bruto", "Mes (1, 2 o 3)", "Estado", "Probabilidad", "Ponderado"], [32, 14, 12, 14, 12, 14])
    for f in range(5, 35):
        celda(ws, f"A{f}", entrada=True)
        celda(ws, f"B{f}", entrada=True, fmt=NUM)
        celda(ws, f"C{f}", entrada=True, fmt='0')
        celda(ws, f"D{f}", entrada=True)
        celda(ws, f"E{f}", f'=IF(D{f}="","",IF(D{f}="Confirmado",1,IF(D{f}="Probable",0.7,0.3)))', fmt='0%')
        celda(ws, f"F{f}", f'=IF(OR(B{f}="",E{f}=""),"",B{f}*E{f})', fmt=NUM)
    lista(ws, "C5:C34", ["1", "2", "3"])
    lista(ws, "D5:D34", ["Confirmado", "Probable", "Posible"])
    ws.column_dimensions["G"].width = 3
    ws.column_dimensions["H"].width = 36
    for col in "IJK":
        ws.column_dimensions[col].width = 13
    encabezado(ws, 4, 8, ["Resumen por mes", "Mes 1", "Mes 2", "Mes 3"])
    filas = [
        ("Bruto ponderado", lambda m: f"=SUMIF($C$5:$C$34,{m},$F$5:$F$34)"),
        ("− Reserva de impuestos", lambda m: f"=-{chr(72 + m)}5*'3 Sistema'!$C$6"),
        ("− Gastos de la actividad", lambda m: "=-'3 Sistema'!$C$8"),
        ("Disponible esperado", lambda m: f"=SUM({chr(72 + m)}5:{chr(72 + m)}7)"),
        ("Sueldo meta", lambda m: "='3 Sistema'!$C$15"),
        ("Diferencia", lambda m: f"={chr(72 + m)}8-{chr(72 + m)}9"),
        ("Señal", lambda m: f'=IF({chr(72 + m)}9=0,"",IF({chr(72 + m)}10<0,"Mes flojo: trabajar","OK"))'),
    ]
    for i, (t, fn) in enumerate(filas, start=5):
        etiqueta(ws, f"H{i}", t, negrita=t in ("Disponible esperado", "Señal"))
        for m in (1, 2, 3):
            celda(ws, f"{chr(72 + m)}{i}", fn(m), fmt=NUM if t != "Señal" else None, negrita=t in ("Disponible esperado", "Señal"))
    ws.conditional_formatting.add("I11:K11", CellIsRule(operator="equal", formula=['"Mes flojo: trabajar"'], fill=PatternFill("solid", fgColor="FBD9BD")))
    return ws


# ----------------------------------------------------------- caso Valeria
def cargar_ejemplo(wb):
    ing = wb["1 Ingresos"]
    for i, v in enumerate([2900, 1800, 4200, 2400, 3600, 1500, 2700, 3900, 2100, 5200, 2600, 1900]):
        ing[f"B{5 + i}"] = v
    clientes = [("Cliente A (agencia)", 13224), ("Cliente B (software)", 7656), ("Cliente C (consultora)", 5220),
                ("Cliente D", 1900), ("Cliente E", 1600), ("Cliente F", 1500), ("Cliente G", 1400), ("Cliente H", 1200), ("Cliente I", 1100)]
    for i, (c, v) in enumerate(clientes):
        ing[f"A{23 + i}"] = c
        ing[f"B{23 + i}"] = v

    g = wb["2 Gastos"]
    gastos = [
        ("Alquiler", "Esencial", 600, "Sí"), ("Servicios (luz, agua, gas)", "Esencial", 80, "Sí"),
        ("Supermercado (incluye mascota)", "Esencial", 300, "No"), ("Transporte", "Esencial", 50, "No"),
        ("Seguro o cobertura de salud", "Esencial", 90, "Sí"), ("Teléfono e internet del hogar", "Esencial", 40, "Sí"),
        ("Gimnasio", "Importante", 30, "Sí"), ("Cursos y formación", "Importante", 40, "No"),
        ("Ropa", "Importante", 40, "No"), ("Veterinario y extras de la mascota", "Importante", 35, "No"),
        ("Restaurantes, salidas y delivery", "Flexible", 120, "No"), ("Suscripciones de entretenimiento", "Flexible", 25, "Sí"),
    ]
    for i, (a, b, c, d) in enumerate(gastos, start=6):
        g[f"A{i}"], g[f"B{i}"], g[f"C{i}"], g[f"D{i}"] = a, b, c, d

    s = wb["3 Sistema"]
    s["C5"], s["C6"], s["C7"], s["C8"] = "USD", 0.20, "Viernes / día 1", 200
    s["C29"], s["C30"], s["C31"] = "No", "No", "No"

    c = wb["4 Cobros"]
    cobros = [(dt.date(2026, 3, 4), "Cliente A", 1800), (dt.date(2026, 3, 12), "Cliente B", 1500), (dt.date(2026, 3, 25), "Cliente D", 900)]
    for i, (fecha, cli, monto) in enumerate(cobros, start=5):
        c[f"A{i}"], c[f"B{i}"], c[f"C{i}"], c[f"I{i}"] = fecha, cli, monto, "Sí"

    e = wb["5 Embalse"]
    antes = [2120, 2160, 4120, 4240, 5320, 4720, 5080, 6400, 6280, 8640, 8920, 8640]
    for i, v in enumerate(antes):
        e[f"A{5 + i}"] = dt.date(2026 + (1 + i) // 12, (1 + i) % 12 + 1, 1)  # 1 feb 2026 … 1 ene 2027
        e[f"B{5 + i}"] = v

    p = wb["6 Previsiones"]
    for i, (a, m, v) in enumerate([("Dentista", 3, 120), ("Viaje anual", 7, 1200), ("Anteojos", 9, 120), ("Regalos y fiestas", 12, 300)], start=11):
        p[f"B{i}"], p[f"C{i}"], p[f"D{i}"] = a, m, v

    t = wb["7 Tarifa"]
    t["C6"], t["C19"] = 3600, 35

    k = wb["8 Cartera 90 días"]
    for i, fila in enumerate([("Cliente B · acuerdo mensual", 648, 1, "Confirmado"), ("Cliente A · rediseño app", 2400, 1, "Probable"),
                              ("Cliente C · mantenimiento", 360, 2, "Confirmado"), ("Cliente nuevo · landing", 1200, 2, "Posible"),
                              ("Cliente B · acuerdo mensual", 648, 2, "Confirmado"), ("Cliente B · acuerdo mensual", 648, 3, "Confirmado"),
                              ("Cliente C · mantenimiento", 360, 3, "Confirmado"), ("Cliente A · fase 2", 1800, 3, "Posible")], start=5):
        k[f"A{i}"], k[f"B{i}"], k[f"C{i}"], k[f"D{i}"] = fila


def construir(ejemplo: bool) -> Path:
    wb = Workbook()
    hoja_inicio(wb)
    hoja_ingresos(wb)
    hoja_gastos(wb)
    hoja_sistema(wb)
    hoja_cobros(wb)
    hoja_embalse(wb)
    hoja_previsiones(wb)
    hoja_tarifa(wb)
    hoja_cartera(wb)
    if ejemplo:
        cargar_ejemplo(wb)
        wb["Inicio"]["B2"] = "EJEMPLO con el caso hipotético de Valeria (capítulos 2 a 11). Usa la versión en blanco para tus datos."
        wb["Inicio"]["B2"].font = Font(name="Calibri", size=11, bold=True, color="C8650F")
    for ws in wb.worksheets:
        ws.sheet_properties.tabColor = AGUA if ws.title[0].isdigit() else PROFUNDO
    SALIDA.mkdir(exist_ok=True)
    nombre = "Plantilla-Metodo-Embalse-EJEMPLO.xlsx" if ejemplo else "Plantilla-Metodo-Embalse.xlsx"
    ruta = SALIDA / nombre
    wb.save(ruta)
    return ruta


if __name__ == "__main__":
    print(construir(False))
    print(construir(True))
