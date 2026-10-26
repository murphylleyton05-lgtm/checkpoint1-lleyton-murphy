#!/usr/bin/env python3
"""Genera las imágenes comerciales (mockups) y los bonos en PDF.

Salidas en entregables/:
  imagenes/01-portada.png ... 06-pin-pinterest.png
  Bono-Mapa-del-Metodo-Embalse.pdf      (1 página; también sirve como imán de suscriptores)
  Bono-Cuaderno-de-trabajo.pdf          (anexos A y B del libro, listos para imprimir)
Requiere haber corrido antes construir_ebook.py y crear_plantilla.py.
"""
import json
import shutil
import subprocess
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader, PdfWriter

RAIZ = Path(__file__).resolve().parent.parent
FUENTE = RAIZ / "fuente"
TMP = RAIZ / ".construccion" / "mockups"
ENTREG = RAIZ / "entregables"
IMG = ENTREG / "imagenes"
CAPTURA = RAIZ / "herramientas" / "captura.cjs"
LIBRO = ENTREG / "Ingresos-variables-sueldo-fijo.pdf"

CSS_BASE = f"""
<link rel="stylesheet" href="{(FUENTE / 'estilos.css').as_uri()}">
<style>
  html, body {{ margin:0; padding:0; }}
  body {{ font-family: Inter, Arial, sans-serif; color:#15232b; }}
  * {{ box-sizing:border-box; }}
</style>
"""


def captura(html: Path, salida: Path, ancho: int, alto: int, escala: int = 2):
    subprocess.run(["node", str(CAPTURA), str(html), str(salida), str(ancho), str(alto), str(escala)], check=True)


def pagina_con(texto: str, lector: PdfReader, desde: int = 4) -> int:
    for i in range(desde, len(lector.pages)):
        if texto.lower() in (lector.pages[i].extract_text() or "").lower():
            return i + 1
    raise SystemExit(f"No encontré la página con: {texto}")


def render_pagina(n: int, nombre: str, dpi: int = 110) -> Path:
    base = TMP / nombre
    subprocess.run(["pdftoppm", "-r", str(dpi), "-png", "-singlefile", "-f", str(n), "-l", str(n), str(LIBRO), str(base)], check=True)
    return base.with_suffix(".png")


def valores_plantilla():
    """Recalcula el ejemplo con LibreOffice y lee la ficha del sistema."""
    origen = ENTREG / "Plantilla-Metodo-Embalse-EJEMPLO.xlsx"
    tmpx = TMP / "xlsx"
    tmpx.mkdir(exist_ok=True)
    shutil.copy(origen, tmpx / origen.name)
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", str(tmpx / "r"), str(tmpx / origen.name)],
                   check=True, capture_output=True, timeout=180)
    wb = load_workbook(tmpx / "r" / origen.name, data_only=True)
    ws = wb["3 Sistema"]
    filas = []
    for r in range(4, 46):
        b, c = ws[f"B{r}"].value, ws[f"C{r}"].value
        filas.append((r, b, c, ws[f"B{r}"].fill.fgColor.rgb if ws[f"B{r}"].fill else None, ws[f"C{r}"].fill.fgColor.rgb))
    return filas


def fmt(v):
    if isinstance(v, float) and v < 1 and v > 0:
        return f"{v:.0%}".replace(".", ",")
    if isinstance(v, (int, float)):
        return f"{v:,.0f}".replace(",", ".")
    return "" if v is None else str(v)


def main():
    TMP.mkdir(parents=True, exist_ok=True)
    IMG.mkdir(parents=True, exist_ok=True)
    lector = PdfReader(str(LIBRO))
    mapa = json.loads((RAIZ / ".construccion" / "mapa-paginas.json").read_text())

    # 1 · Portada
    captura(FUENTE / "portada.html", IMG / "01-portada.png", 794, 1123, 2)
    portada = (IMG / "01-portada.png").as_uri()

    # Páginas interiores
    p_diag = render_pagina(pagina_con("Figura 4.1", lector), "int-diagrama")
    p_sueldo = render_pagina(pagina_con("Figura 6.1", lector), "int-sueldo")
    p_sema = render_pagina(pagina_con("justo después de pagarte", lector), "int-semaforo")
    p_ficha = render_pagina(pagina_con("A.1 · La ficha de tu sistema", lector), "int-ficha")
    p_plan = render_pagina(mapa["cap-15"] + 2, "int-plan")
    p_ej = render_pagina(pagina_con("Ejercicio 2.2", lector), "int-ejercicio")

    # Hoja de cálculo simulada con los valores reales del ejemplo
    filas = valores_plantilla()
    trs = []
    for r, b, c, fb, fc in filas:
        es_seccion = b and c is None and fb and fb.endswith("0E3A4A")
        if es_seccion:
            trs.append(f'<tr class="sec"><td class="rn">{r}</td><td colspan="2">{b}</td></tr>')
        elif b:
            cls = "in" if (fc or "").endswith("FFF4D6") else "out"
            negrita = " b" if b.isupper() or b.startswith(("SUELDO", "META", "M (")) else ""
            trs.append(f'<tr><td class="rn">{r}</td><td class="lab{negrita}">{b}</td><td class="val {cls}{negrita}">{fmt(c)}</td></tr>')
    hoja_html = "\n".join(trs[:23])

    plantilla = TMP / "plantilla.html"
    plantilla.write_text(f"""<!doctype html><html lang="es"><head><meta charset="utf-8">{CSS_BASE}
<style>
 body {{ width:1024px; height:1024px; background:linear-gradient(160deg,#eef5f8,#dcebf3); display:flex; align-items:center; justify-content:center; }}
 .ventana {{ width:900px; background:#fff; border-radius:16px; box-shadow:0 30px 70px rgba(14,58,74,.25); overflow:hidden; }}
 .barra {{ background:#0e3a4a; color:#fff; padding:14px 20px; font-size:15px; display:flex; gap:10px; align-items:center; }}
 .barra i {{ width:12px; height:12px; border-radius:50%; background:#f2b77a; display:inline-block; }}
 .barra b {{ margin-left:8px; font-weight:600; }}
 .pest {{ display:flex; gap:2px; background:#e8eef0; padding:0 12px; font-size:12px; }}
 .pest span {{ padding:8px 12px; color:#3d4d55; }} .pest span.act {{ background:#fff; color:#0e3a4a; font-weight:700; border-top:3px solid #1a7fb5; }}
 table {{ width:100%; border-collapse:collapse; font-size:13px; margin:0; }}
 td {{ border-bottom:1px solid #e3e9ec; padding:6px 10px; }}
 td.rn {{ width:36px; color:#8a9aa2; background:#f4f7f8; text-align:center; font-size:11px; }}
 td.lab {{ width:540px; }} td.b {{ font-weight:700; color:#0e3a4a; }}
 td.val {{ text-align:right; font-variant-numeric:tabular-nums; }}
 td.in {{ background:#fff4d6; }} td.out {{ background:#e9f3f9; }}
 tr.sec td {{ background:#0e3a4a; color:#fff; font-weight:700; }} tr.sec td.rn {{ background:#0e3a4a; color:#9fd0ea; }}
 .pie {{ padding:12px 20px; font-size:12px; color:#5f6f77; background:#f7f3ec; }}
</style></head><body>
<div class="ventana">
  <div class="barra"><i></i><i style="background:#9fd0ea"></i><i style="background:#cfe4f1"></i><b>Plantilla-Metodo-Embalse.xlsx</b></div>
  <div class="pest"><span>1 Ingresos</span><span>2 Gastos</span><span class="act">3 Sistema</span><span>4 Cobros</span><span>5 Embalse</span><span>6 Previsiones</span><span>7 Tarifa</span><span>8 Cartera</span></div>
  <table>{hoja_html}</table>
  <div class="pie">Ejemplo con el caso hipotético de Valeria · Celdas amarillas: las completas tú · Celdas celestes: se calculan solas · Excel y Google Sheets</div>
</div></body></html>""", encoding="utf-8")
    captura(plantilla, IMG / "05-plantilla.png", 1024, 1024, 2)

    # 2 · Mockup principal
    principal = TMP / "principal.html"
    principal.write_text(f"""<!doctype html><html lang="es"><head><meta charset="utf-8">{CSS_BASE}
<style>
 body {{ width:1024px; height:1024px; overflow:hidden; background:radial-gradient(900px 700px at 30% 20%, #fbf6ee, #efe6d6 70%, #e6dac6); position:relative; }}
 .tablet {{ position:absolute; left:470px; top:150px; width:470px; height:640px; background:#1d2328; border-radius:34px; padding:18px; box-shadow:0 40px 80px rgba(0,0,0,.25); transform:rotate(4deg); }}
 .tablet img {{ width:100%; height:100%; object-fit:cover; object-position:top; border-radius:18px; display:block; }}
 .libro {{ position:absolute; left:90px; top:170px; width:440px; height:622px; perspective:1600px; }}
 .tapa {{ position:absolute; inset:0; transform:rotateY(-24deg); transform-origin:left center; box-shadow:30px 34px 60px rgba(0,0,0,.35); border-radius:3px 8px 8px 3px; overflow:hidden; }}
 .tapa img {{ width:100%; height:100%; display:block; }}
 .tapa:after {{ content:""; position:absolute; inset:0; background:linear-gradient(90deg, rgba(0,0,0,.35) 0, rgba(255,255,255,.12) 2.5%, rgba(0,0,0,0) 7%); }}
 .lomo {{ position:absolute; left:-28px; top:6px; width:30px; height:610px; background:linear-gradient(90deg,#062029,#0b2f3c); transform:rotateY(60deg); transform-origin:right center; }}
 .sello {{ position:absolute; left:84px; bottom:70px; background:#0e3a4a; color:#fbf6ee; font-weight:700; font-size:20px; padding:14px 22px; border-radius:40px; box-shadow:0 10px 24px rgba(14,58,74,.3); letter-spacing:.02em; }}
 .sello span {{ color:#f2b77a; }}
</style></head><body>
 <div class="tablet"><img src="{p_sueldo.as_uri()}"></div>
 <div class="libro"><div class="lomo"></div><div class="tapa"><img src="{portada}"></div></div>
 <div class="sello">Ebook PDF <span>+</span> plantilla Excel / Google Sheets</div>
</body></html>""", encoding="utf-8")
    captura(principal, IMG / "02-mockup-principal.png", 1024, 1024, 2)

    # 3 · Qué incluye
    incluye = TMP / "incluye.html"
    tarjetas = [
        (portada, "El libro completo", "128 páginas · 17 capítulos · 3 casos completos · plan de 30 días"),
        ((IMG / "05-plantilla.png").as_uri(), "Plantilla Método Embalse", "Excel y Google Sheets · versión en blanco y ejemplo resuelto"),
        (p_ficha.as_uri(), "Cuaderno de trabajo", "Hojas imprimibles y checklists para usar todos los meses"),
        (p_diag.as_uri(), "Mapa del método", "Todo el sistema en una página, para tener a la vista"),
    ]
    cards = "".join(f'<div class="card"><div class="img"><img src="{u}"></div><b>{t}</b><p>{d}</p></div>' for u, t, d in tarjetas)
    incluye.write_text(f"""<!doctype html><html lang="es"><head><meta charset="utf-8">{CSS_BASE}
<style>
 body {{ width:1024px; height:1024px; background:#0e3a4a; color:#fbf6ee; padding:64px 60px; }}
 h1 {{ font-family:Fraunces, serif; color:#fbf6ee; font-size:58px; margin:0 0 6px; letter-spacing:-.02em; }}
 h1 em {{ color:#f2b77a; font-style:italic; font-weight:400; }}
 .sub {{ color:#9fd0ea; font-size:20px; margin-bottom:38px; }}
 .grid {{ display:grid; grid-template-columns:1fr 1fr; gap:26px; }}
 .card {{ background:#fbf6ee; color:#15232b; border-radius:18px; padding:18px 18px 16px; }}
 .img {{ height:250px; border-radius:10px; overflow:hidden; background:#e9f3f9; margin-bottom:14px; }}
 .img img {{ width:100%; height:100%; object-fit:cover; object-position:top; display:block; }}
 .card b {{ font-size:22px; color:#0e3a4a; display:block; margin-bottom:4px; }}
 .card p {{ margin:0; font-size:15px; color:#3d4d55; line-height:1.4; }}
</style></head><body>
 <h1>Qué <em>incluye</em></h1>
 <div class="sub">Un sistema completo, no solo un libro</div>
 <div class="grid">{cards}</div>
</body></html>""", encoding="utf-8")
    captura(incluye, IMG / "03-que-incluye.png", 1024, 1024, 2)

    # 4 · Interior
    interior = TMP / "interior.html"
    paginas = [(p_diag, "El sistema en una página"), (p_sema, "Un semáforo que decide por ti"), (p_ej, "Ejercicios con tus números"), (p_plan, "Plan de 30 días")]
    figs = "".join(f'<figure><img src="{p.as_uri()}" style="object-position:{"50% 80%" if "semáforo" in t else "top"}"><figcaption>{t}</figcaption></figure>' for p, t in paginas)
    interior.write_text(f"""<!doctype html><html lang="es"><head><meta charset="utf-8">{CSS_BASE}
<style>
 body {{ width:1024px; height:1024px; background:#f7f3ec; padding:48px 52px; }}
 h1 {{ font-family:Fraunces, serif; font-size:46px; color:#0e3a4a; margin:0 0 26px; letter-spacing:-.02em; }}
 .grid {{ display:grid; grid-template-columns:1fr 1fr; gap:26px 30px; }}
 figure {{ margin:0; }}
 figure img {{ width:100%; height:372px; object-fit:cover; object-position:top; border-radius:8px; box-shadow:0 14px 30px rgba(14,58,74,.18); display:block; background:#fff; }}
 figcaption {{ font-size:17px; font-weight:600; color:#0e3a4a; margin-top:10px; }}
</style></head><body>
 <h1>Mira por dentro</h1>
 <div class="grid">{figs}</div>
</body></html>""", encoding="utf-8")
    captura(interior, IMG / "04-interior.png", 1024, 1024, 2)

    # 6 · Pin de Pinterest (1000 x 1500)
    pin = TMP / "pin.html"
    pin.write_text(f"""<!doctype html><html lang="es"><head><meta charset="utf-8">{CSS_BASE}
<style>
 body {{ width:500px; height:750px; background:#fbf6ee; padding:34px 32px; position:relative; overflow:hidden; }}
 .tag {{ font-size:12px; font-weight:700; letter-spacing:.16em; color:#1a7fb5; text-transform:uppercase; }}
 h1 {{ font-family:Fraunces, serif; font-size:40px; line-height:1.02; color:#0e3a4a; margin:10px 0 14px; letter-spacing:-.02em; }}
 h1 em {{ color:#c8650f; font-style:italic; font-weight:400; }}
 ol {{ list-style:none; padding:0; margin:0; counter-reset:p; }}
 li {{ counter-increment:p; position:relative; padding:10px 0 10px 46px; border-bottom:1px solid #e7dfd1; font-size:15.5px; line-height:1.35; }}
 li:before {{ content:counter(p); position:absolute; left:0; top:9px; width:32px; height:32px; border-radius:50%; background:#0e3a4a; color:#fff; font-weight:700; display:flex; align-items:center; justify-content:center; font-size:15px; }}
 li b {{ color:#0e3a4a; }}
 .libro {{ position:absolute; right:-30px; bottom:-40px; width:210px; transform:rotate(-8deg); box-shadow:0 20px 40px rgba(0,0,0,.25); }}
 .pie {{ position:absolute; left:32px; bottom:28px; font-size:13px; color:#3d4d55; width:250px; line-height:1.4; }}
 .pie b {{ color:#0e3a4a; }}
</style></head><body>
 <div class="tag">Freelancers · independientes</div>
 <h1>¿Cobras distinto cada mes? <em>Págate un sueldo fijo.</em></h1>
 <ol>
  <li><b>Separa impuestos</b> de cada cobro, el mismo día.</li>
  <li><b>Junta todo en un «embalse»</b>, una cuenta aparte.</li>
  <li><b>Págate siempre lo mismo</b>, el mismo día del mes.</li>
  <li><b>Sigue el semáforo:</b> verde, amarillo, naranja o rojo.</li>
 </ol>
 <div class="pie"><b>Método Embalse</b><br>Guía + plantilla para ingresos variables</div>
 <img class="libro" src="{portada}">
</body></html>""", encoding="utf-8")
    captura(pin, IMG / "06-pin-pinterest.png", 500, 750, 2)

    # Bono: Mapa del método (A4, 1 página)
    diagrama = (FUENTE / "partes" / "20-parte2.html").read_text(encoding="utf-8")
    svg = diagrama[diagrama.index("<svg"):diagrama.index("</svg>") + 6]
    mapa_html = TMP / "mapa.html"
    mapa_html.write_text(f"""<!doctype html><html lang="es"><head><meta charset="utf-8">{CSS_BASE}
<style>
 @page {{ size:A4; margin:0; }}
 body {{ width:210mm; height:297mm; padding:16mm 16mm 12mm; font-family:Inter, Arial, sans-serif; }}
 .tag {{ font-size:8.5pt; font-weight:700; letter-spacing:.18em; color:#1a7fb5; text-transform:uppercase; }}
 h1 {{ font-family:Fraunces, serif; font-size:26pt; color:#0e3a4a; margin:2mm 0 1mm; letter-spacing:-.01em; }}
 .sub {{ font-family:'Source Serif 4', serif; font-style:italic; color:#3d4d55; font-size:11pt; margin-bottom:4mm; }}
 svg {{ width:100%; height:auto; display:block; margin:0 0 3mm; }}
 .grid {{ display:grid; grid-template-columns:1fr 1fr; gap:4mm; font-size:8.7pt; line-height:1.42; }}
 .box {{ border:1px solid #d8e1e5; border-radius:3mm; padding:3mm 3.6mm; }}
 .box b.t {{ display:block; color:#0e3a4a; font-size:10pt; margin-bottom:1.5mm; }}
 ol, ul {{ margin:0; padding-left:4.5mm; }} li {{ margin:.6mm 0; }}
 .sem {{ display:grid; grid-template-columns:repeat(4,1fr); gap:2mm; margin-top:3mm; font-size:8.2pt; }}
 .sem div {{ border-radius:2mm; padding:2mm 2.4mm; }}
 .pie {{ margin-top:3mm; font-size:7.6pt; color:#5f6f77; border-top:1px solid #d8e1e5; padding-top:2mm; display:flex; justify-content:space-between; }}
</style></head><body>
 <div class="tag">Método Embalse · Mapa en una página</div>
 <h1>Ingresos variables, sueldo fijo.</h1>
 <div class="sub">Cinco espacios, cinco reglas y un día de pago.</div>
 {svg}
 <div class="grid">
  <div class="box"><b class="t">Las 5 reglas</b><ol>
   <li><b>Entrada:</b> todo cobro entra a la cuenta de Entrada.</li>
   <li><b>Impuestos:</b> de cada cobro separas tu % el mismo día.</li>
   <li><b>Barrido:</b> cada semana, lo que excede un mes de gastos de la actividad pasa al Embalse.</li>
   <li><b>Compuerta:</b> el día de pago sale tu sueldo fijo. Nunca más, nunca antes.</li>
   <li><b>Rebalse:</b> si el Embalse supera su meta, 10 % disfrute y 90 % a la cascada.</li></ol></div>
  <div class="box"><b class="t">Tus números</b><ul>
   <li><b>Sueldo meta</b> = costo de vida normal, y ≤ 80 % de tu ingreso disponible promedio.</li>
   <li><b>Sueldo de arranque</b> = el mayor entre supervivencia y piso.</li>
   <li><b>Transición:</b> sueldo meta cuando el embalse tenga 2 sueldos antes de pagarte.</li>
   <li><b>Meta del embalse</b> = M × sueldo meta (M: 3 a 12 según tu variabilidad y riesgos).</li>
   <li><b>Cascada:</b> deudas caras → metas → retiro.</li></ul></div>
 </div>
 <div class="sem">
  <div style="background:#eaf4ea;border-top:3px solid #2e7d32"><b style="color:#2e7d32">Verde · 3+ sueldos</b><br>Normal. Rebalse si superas la meta.</div>
  <div style="background:#fbf4d9;border-top:3px solid #c9a100"><b style="color:#8a6d00">Amarillo · 1 a 3</b><br>Sin compromisos fijos nuevos.</div>
  <div style="background:#fdebdc;border-top:3px solid #d9711c"><b style="color:#b84f00">Naranja · 0,5 a 1</b><br>Modo ajuste + plan comercial.</div>
  <div style="background:#fbe5e5;border-top:3px solid #a61b1b"><b style="color:#a61b1b">Rojo · menos de 0,5</b><br>Supervivencia. Impuestos intactos.</div>
 </div>
 <div class="grid" style="margin-top:3mm;">
  <div class="box"><b class="t">Ritual mensual · 45 minutos</b><ol>
   <li>Págate y aporta a Previsiones.</li>
   <li>Mira el semáforo y aplica su protocolo.</li>
   <li>Aplica la cascada si hubo rebalse.</li>
   <li>Actualiza tu historial de 12 meses.</li>
   <li>Revisa si te alcanzó el sueldo.</li>
   <li>Actualiza tu cartera a 90 días.</li>
   <li>Elige una sola mejora para el mes.</li></ol></div>
  <div class="box"><b class="t">Lo que nunca se hace</b><ul>
   <li>Usar la cuenta de Impuestos como colchón.</li>
   <li>Adelantar el sueldo «por esta vez».</li>
   <li>Subir el sueldo después de un mes excelente (solo en la revisión trimestral, con datos).</li>
   <li>Cobrar fuera de la cuenta de Entrada.</li>
   <li>Invertir el embalse en activos de riesgo.</li>
   <li>Mezclar previsiones con el embalse.</li></ul></div>
 </div>
 <div class="pie"><span>Del libro «Ingresos variables, sueldo fijo», de Lleyton Murphy.<br>Material educativo; no es asesoramiento financiero ni fiscal.</span><span style="text-align:right">Ritual semanal: 10 minutos<br>Revisión trimestral: 2 horas</span></div>
</body></html>""", encoding="utf-8")
    captura(mapa_html, ENTREG / "Bono-Mapa-del-Metodo-Embalse.pdf", 794, 1123, 1)

    # Bono: Cuaderno de trabajo (anexos A y B del libro)
    ini = mapa["anexo-a"] + 1          # +1 por la portada
    fin = mapa["anexo-c"]              # última página del anexo B (índice base 0 de la siguiente sección)
    tapa = (FUENTE / "portada.html").read_text(encoding="utf-8")
    tapa = tapa.replace('href="estilos.css"', f'href="{(FUENTE / "estilos.css").as_uri()}"')
    tapa = tapa.replace("<h1>Ingresos variables,<br><em>sueldo fijo.</em></h1>", "<h1>Cuaderno<br><em>de trabajo.</em></h1>")
    tapa = tapa.replace("El sistema paso a paso para que freelancers e independientes se paguen un sueldo estable, reserven sus impuestos y dejen de vivir de cobro en cobro.",
                        "Las hojas de trabajo y checklists del libro «Ingresos variables, sueldo fijo», listas para imprimir y usar todos los meses.")
    tapa = tapa.replace("<b>Incluye:</b> plantilla de cálculo editable, hojas de trabajo imprimibles, protocolos para meses buenos y malos, guía fiscal orientativa para 6 países y plan de implementación de 30 días.",
                        "<b>Incluye:</b> ficha del sistema, registro de cobros, cierre mensual, seguimiento del embalse, revisión trimestral y checklists de referencia rápida.")
    tapa_html = TMP / "tapa-cuaderno.html"
    tapa_html.write_text(tapa, encoding="utf-8")
    captura(tapa_html, TMP / "tapa-cuaderno.pdf", 794, 1123, 1)
    escritor = PdfWriter()
    escritor.append(str(TMP / "tapa-cuaderno.pdf"))
    for i in range(ini - 1, fin):
        escritor.add_page(lector.pages[i])
    escritor.add_metadata({"/Title": "Cuaderno de trabajo · Ingresos variables, sueldo fijo", "/Author": "Lleyton Murphy"})
    with open(ENTREG / "Bono-Cuaderno-de-trabajo.pdf", "wb") as fh:
        escritor.write(fh)
    print("Cuaderno: páginas", ini, "a", fin, "del libro")


if __name__ == "__main__":
    main()
